"""NFL historical data ingester.

XLSX format with columns: Date, Home Team, Away Team, Home Score, Away Score,
Overtime?, Playoff Game?, Neutral Venue?,
Home Odds Open/Min/Max/Close, Away Odds Open/Min/Max/Close,
Home Line Open/Min/Max/Close, Away Line Open/Min/Max/Close,
Home Line Odds Open/Min/Max/Close, Away Line Odds Open/Min/Max/Close,
Total Score Open/Min/Max/Close,
Total Score Over/Under Open/Min/Max/Close, Notes
"""

from __future__ import annotations

import logging
from pathlib import Path

from bet_agent.db.models import HistoricalMatch, Sport
from bet_agent.ingest.base import BaseIngester, safe_float, safe_int

logger = logging.getLogger(__name__)


def _read_xlsx(file_path: Path) -> list[dict]:
    """Read an XLSX file into a list of dicts using openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    # First row is header
    header = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        d = {}
        for i, val in enumerate(row):
            if i < len(header):
                d[header[i]] = str(val).strip() if val is not None else ""
        result.append(d)
    return result


class NFLIngester(BaseIngester):
    sport = Sport.AMERICAN_FOOTBALL
    source_name = "nfl_historical_xlsx"

    def __init__(self, default_season: str = "unknown") -> None:
        super().__init__()
        self._default_season = default_season

    def parse_file(self, file_path: Path) -> list[dict]:
        rows = _read_xlsx(file_path)
        return [
            r for r in rows
            if r.get("Home Team") and r.get("Away Team") and r.get("Date")
        ]

    def ingest_directory(self, session, dir_path, glob="*.xlsx"):
        return super().ingest_directory(session, dir_path, glob=glob)

    def row_to_model(self, row: dict, source_file: str) -> HistoricalMatch:
        from bet_agent.ingest.base import parse_date

        match_date = parse_date(
            row["Date"],
            formats=["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"],
        )

        season = self._derive_season(source_file, match_date)
        home_score = safe_int(row.get("Home Score"))
        away_score = safe_int(row.get("Away Score"))
        result = "H" if home_score > away_score else ("A" if away_score > home_score else "D")

        # Match metadata
        match_stats: dict = {}
        if row.get("Overtime?") and row["Overtime?"].lower() in ("yes", "true", "1"):
            match_stats["overtime"] = True
        if row.get("Playoff Game?") and row["Playoff Game?"].lower() in ("yes", "true", "1"):
            match_stats["playoff"] = True
        if row.get("Neutral Venue?") and row["Neutral Venue?"].lower() in ("yes", "true", "1"):
            match_stats["neutral_venue"] = True
        if row.get("Notes"):
            match_stats["notes"] = row["Notes"]

        # Moneyline odds (open/min/max/close for each side)
        odds: dict = {}
        for side in ["Home", "Away"]:
            for variant in ["Open", "Min", "Max", "Close"]:
                col = f"{side} Odds {variant}"
                val = safe_float(row.get(col))
                if val is not None:
                    odds[f"{side.lower()}_odds_{variant.lower()}"] = val

        # Spread / line data
        betting_lines: dict = {}
        for side in ["Home", "Away"]:
            for variant in ["Open", "Min", "Max", "Close"]:
                # Point spread
                col = f"{side} Line {variant}"
                val = safe_float(row.get(col))
                if val is not None:
                    betting_lines[f"{side.lower()}_line_{variant.lower()}"] = val

                # Spread odds (juice)
                col = f"{side} Line Odds {variant}"
                val = safe_float(row.get(col))
                if val is not None:
                    betting_lines[f"{side.lower()}_line_odds_{variant.lower()}"] = val

        # Total score lines
        for variant in ["Open", "Min", "Max", "Close"]:
            col = f"Total Score {variant}"
            val = safe_float(row.get(col))
            if val is not None:
                betting_lines[f"total_{variant.lower()}"] = val

            # Over/under odds
            for ou in ["Over", "Under"]:
                col = f"Total Score {ou} {variant}"
                val = safe_float(row.get(col))
                if val is not None:
                    betting_lines[f"total_{ou.lower()}_{variant.lower()}"] = val

        return HistoricalMatch(
            sport=Sport.AMERICAN_FOOTBALL,
            season=season,
            division="NFL",
            match_date=match_date,
            home_team=row["Home Team"].strip(),
            away_team=row["Away Team"].strip(),
            home_score=home_score,
            away_score=away_score,
            result=result,
            match_stats=match_stats,
            odds=odds,
            betting_lines=betting_lines,
            advanced_stats={},
            source=self.source_name,
            source_file=source_file,
        )

    def _derive_season(self, source_file: str, match_date) -> str:
        """NFL season: Aug-Feb spans two calendar years."""
        import re
        name = Path(source_file).stem
        m = re.search(r"(20\d{2})", name)
        if m:
            return m.group(1)
        year = match_date.year
        month = match_date.month
        if month <= 6:
            return str(year - 1)
        return str(year)
